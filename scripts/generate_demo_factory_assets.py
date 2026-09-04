"""Generate synthetic multi-format FACTORY-DEMO-01 document assets reproducibly."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation

PROJECT_ROOT = Path(__file__).resolve().parents[1]
ASSET_ROOT = PROJECT_ROOT / "demo_factory"
DOCUMENTS = (
    (
        "public/Factory_Overview.pdf",
        "Factory Overview",
        "PUBLIC",
        "application/pdf",
        None,
        "FACTORY-DEMO-01, stations S01-S05",
    ),
    (
        "public/S04_Quality_Inspection_Overview.pptx",
        "S04 Quality Inspection Overview",
        "PUBLIC",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "S04",
        "QUALITY-09 inspection decision flow",
    ),
    (
        "internal/Operator_Handbook.docx",
        "Operator Handbook",
        "INTERNAL",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "S02",
        "Positioning warnings and E-STOP-17 escalation",
    ),
    (
        "internal/Station_Layout.xlsx",
        "Station Layout",
        "INTERNAL",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        None,
        "S01 Material Intake through S05 Packaging",
    ),
    (
        "confidential/Positioning_Error_Troubleshooting.pdf",
        "Positioning Error Troubleshooting",
        "CONFIDENTIAL",
        "application/pdf",
        "S02",
        "P4711, QUALITY-09 and repeated positioning error investigation",
    ),
    (
        "confidential/Maintenance_Report_S02.docx",
        "Maintenance Report S02",
        "CONFIDENTIAL",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "S02",
        "encoder replacement, homing, dry cycle, inspection verification",
    ),
    (
        "confidential/Failure_Analysis_S02.pptx",
        "Failure Analysis S02",
        "CONFIDENTIAL",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "S02",
        "P4801 P4802 P4805 P4811 repeated station failure",
    ),
    (
        "restricted/Process_Recipe.xlsx",
        "Process Recipe",
        "RESTRICTED",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "S03",
        "customer-specific assembly force and thermal dwell process recipe",
    ),
    (
        "restricted/Robot_Calibration_Parameters.xlsx",
        "Robot Calibration Parameters",
        "RESTRICTED",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "S03",
        "robot trajectory limits and calibration offsets",
    ),
)
IMAGES = (
    (
        "images/S02_Positioning_Reference.png",
        "S02 Positioning Reference",
        "INTERNAL",
        "S02",
        "positioning reference image",
    ),
    (
        "images/S04_Inspection_Good.png",
        "S04 Inspection Good Reference",
        "PUBLIC",
        "S04",
        "accepted inspection reference image",
    ),
    (
        "images/S04_Inspection_Defect.png",
        "S04 Inspection Defect Reference",
        "CONFIDENTIAL",
        "S04",
        "QUALITY-09 defect reference image",
    ),
)


def main() -> None:
    for relative_path, title, _, _, _, content in DOCUMENTS:
        path = ASSET_ROOT / "documents" / relative_path
        path.parent.mkdir(parents=True, exist_ok=True)
        _write_document(path, title, content)
    for relative_path, title, _, _, label in IMAGES:
        path = ASSET_ROOT / relative_path
        path.parent.mkdir(parents=True, exist_ok=True)
        _write_image(path, label)
    metadata = [
        {
            "document_id": _stable_document_id(relative_path),
            "title": title,
            "classification": classification,
            "mime_type": mime_type,
            "source_system": "FACTORY-DEMO-01-DOCUMENTS",
            "factory_code": "FACTORY-DEMO-01",
            "station_code": station_code,
            "version": "1.0",
            "valid_from": "2026-01-01",
            "tags": content.split(", "),
            "file_path": f"documents/{relative_path}",
            "checksum": _checksum(ASSET_ROOT / "documents" / relative_path),
        }
        for relative_path, title, classification, mime_type, station_code, content in DOCUMENTS
    ]
    metadata.extend(
        {
            "document_id": _stable_document_id(relative_path),
            "title": title,
            "classification": classification,
            "mime_type": "image/png",
            "source_system": "FACTORY-DEMO-01-DOCUMENTS",
            "factory_code": "FACTORY-DEMO-01",
            "station_code": station_code,
            "version": "1.0",
            "valid_from": "2026-01-01",
            "tags": [label],
            "file_path": relative_path,
            "checksum": _checksum(ASSET_ROOT / relative_path),
        }
        for relative_path, title, classification, station_code, label in IMAGES
    )
    metadata_path = ASSET_ROOT / "metadata" / "document_catalog.json"
    metadata_path.parent.mkdir(parents=True, exist_ok=True)
    metadata_path.write_text(json.dumps(metadata, indent=2) + "\n", encoding="utf-8")


def _write_document(path: Path, title: str, content: str) -> None:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        _write_pdf(path, title, content)
    elif suffix == ".docx":
        document = Document()
        document.add_heading(title, level=0)
        document.add_paragraph(content)
        document.add_paragraph("Synthetic training data for FACTORY-DEMO-01 only.")
        document.save(path)
    elif suffix == ".pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = content
        presentation.save(path)
    elif suffix == ".xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "FACTORY-DEMO-01"
        sheet.append(["Title", title])
        sheet.append(["Content", content])
        sheet.append(["Notice", "Synthetic training data only"])
        workbook.save(path)
    else:
        raise ValueError(f"Unsupported document type: {path}")


def _write_pdf(path: Path, title: str, content: str) -> None:
    escaped = (
        (title + " - " + content)
        .replace("\\", "\\\\")
        .replace("(", "\\(")
        .replace(")", "\\)")
    )
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1", "replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(value)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    path.write_bytes(output)


def _write_image(path: Path, label: str) -> None:
    image = Image.new("RGB", (960, 540), color="#e8f1ef")
    draw = ImageDraw.Draw(image)
    draw.rectangle((48, 48, 912, 492), outline="#176b4b", width=8)
    draw.text((88, 120), "FACTORY-DEMO-01", fill="#176b4b")
    draw.text((88, 220), label, fill="#17211f")
    image.save(path, format="PNG")


def _stable_document_id(relative_path: str) -> str:
    return "doc-" + hashlib.sha256(relative_path.encode()).hexdigest()[:16]


def _checksum(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


if __name__ == "__main__":
    main()
